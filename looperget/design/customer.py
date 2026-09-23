# -*- coding: utf-8 -*-
"""
looperget.design.customer — 고객 전달본 (V117 2단계 · 2-B 두 벌 · 2-C 연락처 줄 · 2-D 반경 표기).

발행 한 번에 두 벌을 낸다(publish.run):
  · 내부 검토본 = 지금 출력 그대로(렌더러 무변경) — 파일명 끝 `_내부검토`
  · 고객 전달본 = 내부본을 **이 모듈이 손질한 사본** — 기존 파일명(30_제안서… · 40_견적서… · 41_…)
손질은 **삭제·이동·치환만**이다(새 고객 문구 없음 · 불변 원칙 10). 무엇을 어떻게 = `customer_text.json` 한 파일.
  a. 표지 관문 도장 제거                     b. 대괄호 태그 → 본문에서 떼고 「고객께서 확인해 주실 것」 목록으로
  c. 규칙 번호·「대표 지정」·「(분배점 파생)」 제거   d. 품목·세트 코드 → 몰 상품명(없으면 코드만 삭제)
  e. 엔진 설명 부록 면 제거                   f. 엔진 설명 문장 → 삭제 또는 짧은 라벨
  2-C 연락처·회신 한 줄(값 또는 밑줄)          2-D 설계 반경 = 「당사 보수 기준」 라벨 · 0.5 m 내림(R·Ø 같은 값에서)
새 문구가 필요한 자리 = JSON 의 needs_copy(값 비움 → fallback 처리) — Fable 창에서 채운다.

순수 부분(clean · radius_c · contact_line · scan_text …)은 파일을 읽지 않는다. pptx()/xlsx_summary() 는 입출력 껍데기다.

[V117 · 3차 검토 보완(미배포 중 · PKG 110 유지)] D2 태그 사항 = 괄호 속 「 · 」로 쪼개지 않음 · 괄호 짝이 깨지면 문장 단위 ·
  D3 부록 면은 지우기 **전에** 태그 사항을 모은다(문구는 같은 조건의 관문 항목) · 「주배관 n mm — 기본(승인본)」의 기본값은 확인 목록에 ·
  D7 반경 내림 = 반올림 **전** 엔진 값에서(raw_radii) — 지면 글자 「12」만 보고 내리면 R 11.96 이 12 로 나갔다.
"""
from __future__ import annotations

import copy
import json
import math
import os
import re
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

RULES_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "customer_text.json")
CONTACT_FIELDS = (("name", "담당자"), ("phone", "연락처"), ("email", "이메일"), ("reply_by", "회신 요청일"))
BLANK = "____________"
ASK_TITLE = "고객께서 확인해 주실 것"       # 현장 요약 면의 기존 확인사항 목록(render_pptx._site_summary)
CONTACT_Y, ASK_BOTTOM = 6.72, 6.58          # 연락처 줄 자리(in) · 확인 목록이 내려올 수 있는 한계(in)

_CACHE: Dict[str, Dict] = {}


def load_rules(path: Optional[str] = None) -> Dict:
    p = path or RULES_PATH
    if p not in _CACHE:
        with open(p, encoding="utf-8") as f:
            _CACHE[p] = json.load(f)
    return _CACHE[p]


# ══════════════ 2-D 반경 표기 ══════════════
def _m1(v) -> str:
    return ("%.1f" % float(v)).rstrip("0").rstrip(".")


def radius_c(r: float) -> float:
    """고객본 설계 반경 = **0.5 m 단위 내림**(당사 보수 기준 · 커버를 크게 말하지 않는 쪽). Ø = 2 × 이 값(= 지름 정수 m 내림).
    r = **반올림 전 엔진 값**이어야 한다(D7) — 지면 글자는 raw_radii 로 원값을 찾아 내린다(round_radius)."""
    return math.floor(round(2.0 * float(r), 6)) / 2.0


def raw_radii(S: Optional[Dict]) -> List[float]:
    """[D7] 지면에 찍힌 반경들의 **반올림 전 엔진 값**(순수) — 살수원(설계 말단압)·안쪽 원·구역 말단(hydro).
    hydro 의 p_end 는 소수 2자리로 반올림돼 있다 → 그 구간의 아래끝(−0.005 bar)에서 다시 계산(작게 말하는 쪽)."""
    from . import hydro_zone as HZ
    S = S or {}
    sp = S.get("spray") or {}
    model = sp.get("model")
    out: List[float] = []
    try:
        if sp.get("p_bar") is not None:
            r = float(HZ.radius_m(float(sp["p_bar"]), model))
            out.append(r)
            out.append(min(float(HZ.head_profile(model).get("r_in", 7.0)), r))
    except Exception:
        pass
    for h in S.get("hydro") or []:
        try:
            if h.get("p_end") is not None:
                out.append(float(HZ.radius_m(max(0.0, float(h["p_end"]) - 0.005), model)))
        except Exception:
            pass
    return out


def _raw_of(v: float, raws: Sequence[float]) -> Optional[float]:
    """지면 값(소수 1자리로 찍힌 것) → 같은 글자로 찍히는 엔진 원값 중 가장 작은 것 · 없으면 None."""
    key = "%.1f" % round(float(v), 1)
    c = [float(x) for x in raws or () if "%.1f" % round(float(x), 1) == key]
    return min(c) if c else None


def _r_cust(v: float, raws: Sequence[float], integer: bool) -> Optional[float]:
    """지면 반경 값 v → 고객본 반경. 원값을 모르면: 소수 표기는 반올림 구간 아래끝(v−0.05)에서 내림 · 정수 표기는 손대지 않음(None)."""
    raw = _raw_of(v, raws)
    if raw is None:
        if integer:
            return None
        raw = float(v) - 0.05
    return radius_c(raw)


def spec_diam(model: Optional[str] = None) -> str:
    """제품 사양 한 줄 = 제조사 성능표(HEAD_PROFILES curve)의 지름 범위 + 압력 조건."""
    from .layout import HEAD_PROFILES
    cv = (HEAD_PROFILES.get(model or "427B") or {}).get("curve") or []
    if not cv:
        return ""
    ps, ds = [float(b) for b, _q, _d in cv], [float(d) for _b, _q, d in cv]
    return "최대 살수 직경 %s~%sm (%s~%s bar)" % (_m1(min(ds)), _m1(max(ds)), _m1(min(ps)), _m1(max(ps)))


_RX_R = re.compile(r"(R |반경 )(\d+(?:\.\d)?)(?= ?m)")
_RX_D = re.compile(r"(Ø )(\d+(?:\.\d)?)(?= ?m)")
_RX_FIG = re.compile(r"^(\d+(?:\.\d)?)m$")


def round_radius(t: str, raws: Sequence[float] = ()) -> str:
    """지면 글자의 R·Ø → 고객본 값. raws = raw_radii(S)(반올림 전 엔진 값) — R 은 원값에서 0.5 m 내림, Ø = 2 × 그 R(K-07).
    Ø 글자 = 2 × (반올림한 R) 이므로 Ø/2 로 같은 원값을 찾는다. 원값을 모르면 _r_cust 규칙(보수 쪽 · 정수는 그대로)."""
    def _r(m):
        v = m.group(2)
        rc = _r_cust(float(v), raws, "." not in v)
        return m.group(0) if rc is None else m.group(1) + _m1(rc)

    def _d(m):
        v = m.group(2)
        rc = _r_cust(float(v) / 2.0, raws, "." not in v)
        return m.group(0) if rc is None else m.group(1) + _m1(2 * rc)

    def _f(m):
        v = m.group(1)
        rc = _r_cust(float(v) / 2.0, raws, "." not in v)
        return m.group(0) if rc is None else _m1(2 * rc) + "m"
    t = _RX_R.sub(_r, t)
    t = _RX_D.sub(_d, t)
    return _RX_FIG.sub(_f, t)


# ══════════════ 2-C 연락처 ══════════════
def contact_of(meta: Optional[Dict]) -> Dict[str, str]:
    c = (meta or {}).get("contact") or {}
    return {k: str(c.get(k) or "").strip() for k, _lab in CONTACT_FIELDS}


def contact_line(c: Dict[str, str], blank: str = BLANK) -> str:
    """한 줄 묶음 — 라벨은 짧은 명사만 · 값이 없으면 손으로 적을 밑줄."""
    return "  ·  ".join("%s %s" % (lab, (c.get(k) or "").strip() or blank) for k, lab in CONTACT_FIELDS)


def contact_manager(c: Dict[str, str], default: str = "") -> str:
    """견적서의 기존 연락처 칸(담당자) — 값이 있으면 「이름 · 연락처 · 이메일」, 없으면 기존 담당자."""
    parts = [c.get(k) for k in ("name", "phone", "email") if c.get(k)]
    return " · ".join(parts) if parts else default


# ══════════════ 이름 사전 ══════════════
def name_map(price_db: Optional[Dict], bom: Sequence[Dict] = ()) -> Dict[str, str]:
    """품목 코드 → 몰 상품명(Products 시트 = price_db · BOM 행). 「단가 DB 미연결」 임시 이름은 쓰지 않는다."""
    out = {}
    for c, v in (price_db or {}).items():
        nm = str((v or {}).get("name") or "").strip() if isinstance(v, dict) else ""
        if nm:
            out[str(c).zfill(5)] = nm
    for b in bom or ():
        nm = str(b.get("name") or "").strip()
        if nm and "단가 DB 미연결" not in nm:
            out.setdefault(str(b.get("code")).zfill(5), nm)
    return out


def set_map(S: Optional[Dict] = None, sets_db: Optional[Dict] = None) -> Dict[str, str]:
    """세트명 → 한 품목 코드(세트가 품목 하나일 때만). 출처 = heads.KITS · 이 설계의 연결부(엔진) · Sets 시트 item_code."""
    from .heads import KITS
    out = {}
    for k, v in KITS.items():
        if re.fullmatch(r"\d{5}", k) and v.get("set"):
            out[v["set"]] = k
    for c in (S or {}).get("connections") or []:
        rcp = c.get("recipe") or {}
        if len(rcp) == 1 and c.get("name"):
            out.setdefault(c["name"], str(next(iter(rcp))).zfill(5))
    for _cat, grp in (sets_db or {}).items():
        for nm, info in (grp or {}).items() if isinstance(grp, dict) else ():
            ic = str((info or {}).get("item_code") or "").strip() if isinstance(info, dict) else ""
            if re.fullmatch(r"\d{4,5}", ic):
                out.setdefault(nm, ic.zfill(5))
    return out


# ══════════════ 글자 손질(순수) ══════════════
class Ctx:
    """손질 한 번의 문맥 — where = "pptx"|"xlsx" · mode = "full"(고객본) | "tags"(시공업체용 41 — 태그만)."""

    def __init__(self, R: Dict, where: str, mode: str = "full", names: Optional[Dict] = None,
                 sets: Optional[Dict] = None, vars: Optional[Dict] = None, codes: Iterable[str] = ()):
        self.R, self.where, self.mode = R, where, mode
        self.names, self.sets, self.vars = dict(names or {}), dict(sets or {}), dict(vars or {})
        self.codes = set(codes) | set(self.names)
        self.moved: List[str] = []
        self.covered: List[str] = []
        self.hits: Dict[str, int] = {}
        self.unknown_tags: List[str] = []

    def add(self, item: str):
        it = re.sub(r"\s+", " ", str(item or "")).strip(" .,:;·—")
        if len(it) >= 2 and it not in self.moved:
            self.moved.append(it)

    def hit(self, k: str, n: int = 1):
        self.hits[k] = self.hits.get(k, 0) + n


_RX: Dict[str, "re.Pattern"] = {}


def _rx(p: str):
    if p not in _RX:
        _RX[p] = re.compile(p, re.M)
    return _RX[p]


_OPEN, _CLOSE = "(（「", ")）」"


def _depth(s: str) -> int:
    d = 0
    for c in s:
        if c in _OPEN:
            d += 1
        elif c in _CLOSE:
            d -= 1
    return d


def _balanced(s: str) -> bool:
    d = 0
    for c in s:
        if c in _OPEN:
            d += 1
        elif c in _CLOSE:
            d -= 1
            if d < 0:
                return False
    return d == 0


def _split_top(t: str, sep: str) -> List[str]:
    """[D2] 괄호 **밖**의 구분자에서만 나눈다 — 「부속 (호스는 T 양쪽 · 파이프는 엘보)」의 「 · 」는 한 사항 안이다."""
    parts, last = [], 0
    for m in re.finditer(sep, t):
        if m.end() > m.start() and _depth(t[:m.start()]) <= 0:
            parts.append(t[last:m.start()])
            last = m.end()
    parts.append(t[last:])
    return parts


def _item_of(text: str, a: int, b: int, tag_re: str) -> str:
    """태그가 붙은 사항의 원래 문구 — 문장 → 「 — 」 조각 → 「 · 」 조각 중 태그가 있는 것(태그·목록 기호 제거).
    [D2] 괄호 속 구분자로는 나누지 않는다 · 결과의 괄호 짝이 깨지면 문장 단위로 되돌린다."""
    t = text[:a] + "\x00" + text[b:]
    levels = []
    for k, sep in enumerate((r"\n|(?<=[.。])\s+", r"\s+—\s+", r"\s+·\s+")):
        parts = re.split(sep, t) if k == 0 else _split_top(t, sep)     # 문장 단위는 예전 그대로
        t = next((p for p in parts if "\x00" in p), t)
        levels.append(t)

    def fin(x: str) -> str:
        x = _rx(tag_re).sub("", x.replace("\x00", ""))
        x = re.sub(r"^\s*(?:\d+\.|※|▷|·|-)\s*", "", x)
        y = x.strip(" .,:;·—()")
        return y if _balanced(y) else x.strip(" .,:;·—")
    out = fin(levels[-1])
    return out if _balanced(out) else fin(levels[0])


def clean(text: str, ctx: Ctx) -> str:
    """한 문단(또는 칸) 글자 → 고객본 글자. 규칙 → 세트 코드 → 태그 → 품목 코드 → 반경 → 정리 순."""
    R, t = ctx.R, str(text)
    if not t.strip():
        return t
    t0 = t
    full = ctx.mode == "full"
    if full:
        for r in R.get("rules") or []:
            if r.get("where") not in ("all", ctx.where):
                continue
            act = r.get("action", "keep")
            val = r.get("value") or ""
            if r.get("needs_copy") and not val:
                act = r.get("fallback", act)
            if act == "keep":
                continue
            rx = _rx(r["re"])
            n = len(rx.findall(t))
            if not n:
                continue
            if act == "delete":
                repl = ""
            else:
                repl = val.format(**ctx.vars) if "{" in val else val
            if act == "move":
                it = r.get("item") or ""
                if "\\" in it:                              # [D3] 항목에 원문 값(\1)을 쓰는 규칙 — 「주배관 50 mm — 기본」
                    for mm in rx.finditer(t):
                        ctx.add(mm.expand(it))
                else:
                    ctx.add(it)
            if r.get("covered_by") and r["covered_by"] not in ctx.covered:
                ctx.covered.append(r["covered_by"])
            t = rx.sub(repl, t)
            ctx.hit(r["id"], n)
        # d. 세트 코드 — 한 품목이면 그 몰 상품명, 아니면 삭제
        def _set(m):
            code = ctx.sets.get(m.group(0))
            ctx.hit("set_code")
            return ctx.names.get(code, "") if code else ""
        t = _rx(R["set_code_re"]).sub(_set, t)
    # b. 대괄호 태그(41 시공업체용도)
    allow = set(R.get("bracket_allow") or [])
    out, pos = [], 0
    for m in _rx(R["tag_re"]).finditer(t):
        word = m.group(1).strip()
        if m.group(0) in allow or word in ("LHC", "LSS"):
            continue
        act = (R.get("tags") or {}).get(word)
        if act is None:
            act = R.get("tag_default", "delete")
            ctx.unknown_tags.append(word)
        if act == "move":
            ctx.add(_item_of(t, m.start(), m.end(), R["tag_re"]))
        ctx.hit("tag:" + word)
        out.append(t[pos:m.start()])
        pos = m.end()
        if t[pos:pos + 1] == " " and (m.start() == 0 or t[m.start() - 1] in (" ", "\n")):
            pos += 1                                # 태그 앞뒤 빈칸이 겹치지 않게(「3. [미확정] 인입관」 → 「3. 인입관」)
    if out:
        t = "".join(out) + t[pos:]
    if full:
        # d. 품목 코드 → 몰 상품명(없으면 코드만 삭제). 이름이 「…세트」면 뒤따르는 「 세트」를 겹쳐 쓰지 않는다.
        def _code(m):
            code = m.group(1)
            if code in ctx.names:
                ctx.hit("code")
                return "\x01" + ctx.names[code] + ("\x02" if ctx.names[code].endswith("세트") else "")
            if code in ctx.codes or code.startswith("0") or code == "99999":
                ctx.hit("code_drop")
                return ""
            return code
        t = _rx(R["code_re"]).sub(_code, t)
        t = t.replace("\x02 세트", "").replace("\x02", "").replace("\x01", "")
        if ctx.where == "pptx":
            t = round_radius(t, ctx.vars.get("radii") or ())   # 2-D — 설계 반경 0.5 m 내림 · Ø 같은 값에서 · [D7] 원값 기준
    if t == t0:
        return t0                                   # 🔴 손대지 않은 글자는 정리도 하지 않는다 — 고정면 빈칸은 화살표 자리다(면23)
    for a, b in R.get("tidy") or []:
        t = _rx(a).sub(b, t)
    return t


# ══════════════ 기계 검사(순수) ══════════════
def scan_text(items: Iterable[Tuple[str, str]], codes: Iterable[str] = (), R: Optional[Dict] = None) -> List[Dict]:
    """(자리, 글자) 목록 → 고객본에 남으면 안 되는 것 [{where, kind, hit, text}]. 태그·규칙 번호·품목 코드·세트 코드·금지어."""
    R = R or load_rules()
    C = R["check"]
    allow = set(R.get("bracket_allow") or [])
    codes = set(codes)
    out = []
    for where, t in items:
        t = str(t)
        for m in _rx(C["tag_re"]).finditer(t):
            if m.group(0) not in allow:
                out.append({"where": where, "kind": "태그", "hit": m.group(0), "text": t[:90]})
        for m in _rx(C["rule_re"]).finditer(t):
            out.append({"where": where, "kind": "규칙 번호", "hit": m.group(0), "text": t[:90]})
        for m in _rx(C["set_code_re"]).finditer(t):
            out.append({"where": where, "kind": "세트 코드", "hit": m.group(0), "text": t[:90]})
        for m in _rx(C["code_re"]).finditer(t):
            c = m.group(1)
            if c in codes or c.startswith("0") or c == "99999":
                out.append({"where": where, "kind": "품목 코드", "hit": c, "text": t[:90]})
        for w in C["words"]:
            if w in t:
                out.append({"where": where, "kind": "금지어", "hit": w, "text": t[:90]})
    return out


def pptx_items(path_or_prs) -> List[Tuple[str, str]]:
    from pptx import Presentation
    prs = Presentation(path_or_prs) if isinstance(path_or_prs, str) else path_or_prs
    out = []

    def walk(i, shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(i, sh.shapes)
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(("면%d %s" % (i, sh.name), sh.text_frame.text))
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            out.append(("면%d %s 표" % (i, sh.name), c.text))
    for i, sl in enumerate(prs.slides, 1):
        walk(i, sl.shapes)
    return out


def xlsx_items(path: str, keep_code_col: bool = True, first_item_row: int = 10) -> List[Tuple[str, str]]:
    """견적서 글자 칸. keep_code_col = 품목정보 칸(B열)의 **마지막 줄 = 품목 코드 칸**은 빼고 본다(농업·건설 프로필 — 코드 칸 보임)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value.strip():
                    continue
                v = c.value
                if keep_code_col and c.column == 2 and c.row >= first_item_row and "\n" in v:
                    head, last = v.rsplit("\n", 1)
                    if re.fullmatch(r"\d{5}", last.strip()):
                        v = head
                out.append(("%s!%s" % (ws.title, c.coordinate), v))
    return out


# ══════════════ PPTX 고객본(입출력 껍데기) ══════════════
def _em_lines(text: str, box_in: float, pt: float, pad: float = 0.20) -> int:
    em = sum(1.0 if ord(c) > 0x2000 else (0.30 if c == " " else 0.52) for c in text)
    return max(1, int(-(-(em * pt / 72.0) // (box_in - pad))))


def _process_tf(tf, ctx: Ctx) -> Tuple[bool, bool]:
    """글상자 한 개 → (바뀜?, 비었나?). 문단 단위로 손질 — 줄바꿈(a:br)이 든 문단은 run 단위로."""
    from pptx.oxml.ns import qn
    changed = False
    had = bool(tf.text.strip())
    for p in list(tf.paragraphs):
        runs = p.runs
        if not runs:
            continue
        orig = "".join(r.text for r in runs)
        if p._p.findall(qn("a:br")):
            for r in runs:
                nt = clean(r.text, ctx)
                if nt != r.text:
                    r.text, changed = nt, True
            new = "".join(r.text for r in runs)
        else:
            new = clean(orig, ctx)
            if new != orig:
                runs[0].text = new
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
                changed = True
        if orig.strip() and not new.strip() and len(tf.paragraphs) > 1:
            p._p.getparent().remove(p._p)
    return changed, had and not tf.text.strip()


def _walk_shapes(shapes, ctx: Ctx, drops: List[str], report: Dict):
    for sh in list(shapes):
        if sh.shape_type == 6:
            _walk_shapes(sh.shapes, ctx, drops, report)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    _ch, empty = _process_tf(c.text_frame, ctx)
                    if empty:                                   # 칸이 비면 「-」(표의 기존 빈칸 표기)
                        p0 = c.text_frame.paragraphs[0]
                        if p0.runs:
                            p0.runs[0].text = "-"
                        else:
                            c.text = "-"
            continue
        if not sh.has_text_frame:
            continue
        before = sh.text_frame.text
        if any(_rx(d).search(before) for d in drops):
            sh._element.getparent().remove(sh._element)
            report["shapes_dropped"] += 1
            continue
        _ch, empty = _process_tf(sh.text_frame, ctx)
        after = sh.text_frame.text
        if empty or (after != before and any(_rx(d).search(after) for d in drops)):
            sh._element.getparent().remove(sh._element)
            report["shapes_dropped"] += 1


def _find_ask(prs):
    from pptx.util import Emu
    for sl in prs.slides:
        lab = next((sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip() == ASK_TITLE), None)
        if lab is None:
            continue
        cand = [sh for sh in sl.shapes if sh.has_text_frame and sh is not lab and sh.top is not None
                and sh.top > lab.top and abs(int(sh.left or 0) - int(lab.left or 0)) < Emu(914400 * 0.2)
                and sh.text_frame.text.strip().startswith("▷")]
        cand.sort(key=lambda sh: int(sh.top))
        return sl, lab, (cand[0] if cand else None)
    return None, None, None


def _append_items(box, items: List[str]) -> Dict:
    """확인 목록 끝에 항목을 붙인다(마지막 문단 복제 — 서식 그대로). 넘치면 글자만 줄인다(10 → 9 → 8.5 pt)."""
    from pptx.util import Emu, Pt
    tf = box.text_frame
    have = [p.text for p in tf.paragraphs]
    added = []
    for it in items:
        if any(it in h or h.replace("▷", "").strip() in it for h in have if h.strip()):
            continue
        last = tf.paragraphs[-1]._p
        new = copy.deepcopy(last)
        last.addnext(new)
        p = tf.paragraphs[-1]
        runs = p.runs
        runs[0].text = "▷ " + it
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        have.append(runs[0].text)
        added.append(it)
    top = Emu(box.top).inches
    w = Emu(box.width).inches
    size = None
    for sz in (None, 9.0, 8.5):
        s = sz or (tf.paragraphs[0].runs[0].font.size.pt if tf.paragraphs[0].runs and tf.paragraphs[0].runs[0].font.size else 10.0)
        lh = s * 1.22 / 72.0 + 2.0 / 72.0
        bottom = top + sum(_em_lines(p.text, w, s) for p in tf.paragraphs) * lh
        size = s
        if sz:
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(sz)
        if bottom <= ASK_BOTTOM:
            break
    return {"added": added, "size": size, "bottom_in": round(bottom, 2), "overflow": bottom > ASK_BOTTOM}


def _add_contact(slide, box, line: str):
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_textbox(Inches(2.10), Inches(CONTACT_Y), Inches(10.6), Inches(0.30))
    tb.name = "고객 연락처 줄"
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = line
    src = box.text_frame.paragraphs[0].runs[0].font if (box is not None and box.text_frame.paragraphs[0].runs) else None
    r.font.size = Pt(10.5)
    if src is not None:
        if src.name:
            r.font.name = src.name
        try:
            if src.color and src.color.type is not None and src.color.rgb is not None:
                r.font.color.rgb = src.color.rgb
        except Exception:
            pass
    return tb


def _delete_slide(prs, slide):
    from pptx.oxml.ns import qn
    for sid in list(prs.slides._sldIdLst):
        rid = sid.get(qn("r:id"))
        if prs.part.rels[rid].target_part is slide.part:
            prs.part.drop_rel(rid)
            prs.slides._sldIdLst.remove(sid)
            return True
    return False


def _shape_texts(shapes) -> List[str]:
    out = []
    for sh in shapes:
        if sh.shape_type == 6:
            out += _shape_texts(sh.shapes)
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text for r in sh.table.rows for c in r.cells if c.text.strip()]
    return out


def slide_move_tags(slide, R: Dict) -> List[str]:
    """[D3] 한 면의 옮길 태그(tags 표에서 move) 낱말 목록(나온 순서 · 중복 포함)."""
    mv = {k for k, v in (R.get("tags") or {}).items() if v == "move"}
    return [m.group(1).strip() for t in _shape_texts(slide.shapes) for m in _rx(R["tag_re"]).finditer(t)
            if m.group(1).strip() in mv]


def gate_moves(gate: Optional[Dict], R: Dict, ctx: "Ctx") -> List[str]:
    """[D3] 관문 조건부 항목 중 옮길 태그가 붙은 것 → 확인 목록 사항(태그 조각 · 규칙 번호·코드는 같은 손질)."""
    mv = {k for k, v in (R.get("tags") or {}).items() if v == "move"}
    out = []
    for g in (gate or {}).get("conditional") or []:
        for m in _rx(R["tag_re"]).finditer(str(g)):
            if m.group(1).strip() in mv:
                sub = Ctx(R, "pptx", names=ctx.names, sets=ctx.sets, codes=ctx.codes, vars=ctx.vars)
                it = clean(_item_of(str(g), m.start(), m.end(), R["tag_re"]), sub)
                it = re.sub(r"\s+", " ", it).strip(" .,:;·—")
                if len(it) >= 2 and it not in out:
                    out.append(it)
                break
    return out


def pptx(src: str, dst: str, S: Dict, meta: Dict, *, R: Optional[Dict] = None, extra_moved: Sequence[str] = (),
         names: Optional[Dict] = None, sets: Optional[Dict] = None) -> Dict:
    """내부 검토본 PPTX(src) → 고객 전달본(dst). → 보고 {moved, covered, hits, dropped …}."""
    from pptx import Presentation
    R = R or load_rules()
    prs = Presentation(src)
    report = {"slides_dropped": [], "shapes_dropped": 0, "stamp_removed": False}
    names = names if names is not None else name_map(meta.get("price_db") if isinstance(meta.get("price_db"), dict) else {}, S.get("bom") or [])
    sets = sets if sets is not None else set_map(S, meta.get("sets_db"))
    ctx = Ctx(R, "pptx", names=names, sets=sets, codes=[str(b.get("code")) for b in S.get("bom") or []],
              vars={"spec_diam": spec_diam((S.get("spray") or {}).get("model")), "radii": raw_radii(S)})
    # a. 표지 관문 도장
    g = S.get("gate") or {}
    stamp = "● " + str(g.get("label") or "")
    first = prs.slides[0]
    for sh in list(first.shapes):
        if sh.has_text_frame and g.get("level") not in (None, "ok") and sh.text_frame.text.strip() == stamp:
            sh._element.getparent().remove(sh._element)
            report["stamp_removed"] = True
    # e. 부록 면 — [D3] 지우기 **전에** 그 면의 옮길 태그(move)를 모은다. 부록 문장은 태그가 조사에 붙어
    #    (「펌프가 [미확정]이라」「부속은 [미확정]입니다」) 떼면 문장이 깨진다 → 사항 문구는 같은 조건에서 나온 **관문 항목**
    #    (「펌프 [미확정] — …」「부속 [미확정] 자리 n곳(…)」)의 태그 조각을 쓴다(기존 문구 · 겹치면 _append_items 가 거른다).
    drop_moves: List[str] = []
    for rule in R.get("slides_drop") or []:
        for sl in list(prs.slides):
            ts = [sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame]
            if any(_rx(rule["title_re"]).search(t) for t in ts):
                drop_moves += slide_move_tags(sl, R)
                if _delete_slide(prs, sl):
                    report["slides_dropped"].append(rule["id"])
    report["dropped_move_tags"] = drop_moves
    gate_items = gate_moves(S.get("gate"), R, ctx) if drop_moves else []
    drops = [d["re"] for d in R.get("shapes_drop") or []]
    for sl in prs.slides:
        _walk_shapes(sl.shapes, ctx, drops, report)
    # b. 태그 사항 → 기존 확인 목록 · 2-C 연락처 줄
    sl, lab, box = _find_ask(prs)
    items = list(ctx.moved) + [x for x in extra_moved if x not in ctx.moved]
    items += [x for x in gate_items if x not in items]
    if box is not None:
        report["ask"] = _append_items(box, items)
        have = box.text_frame.text
        report["covered_missing"] = [c for c in ctx.covered if c not in have]
        _add_contact(sl, box, contact_line(contact_of(meta)))
        report["contact"] = True
    else:
        report["ask"] = {"added": [], "overflow": False, "missing_face": True}
        report["covered_missing"] = list(ctx.covered)
        report["contact"] = False
    report.update({"moved": items, "covered": ctx.covered, "hits": ctx.hits, "unknown_tags": sorted(set(ctx.unknown_tags))})
    prs.save(dst)
    return report


# ══════════════ XLSX 고객본 — 요약 사본 손질(순수) ══════════════
def xlsx_summary(S: Dict, remarks: str, *, mode: str = "full", R: Optional[Dict] = None,
                 names: Optional[Dict] = None, sets: Optional[Dict] = None) -> Tuple[Dict, str, Ctx]:
    """견적서가 읽는 요약(S)의 품목 이름·규격·비고와 특약사항 → 고객본 사본. 수량·단가·합계는 **그대로**(두 벌 합계 동일)."""
    R = R or load_rules()
    S2 = copy.deepcopy(S)
    ctx = Ctx(R, "xlsx", mode=mode, names=names if names is not None else name_map({}, S.get("bom") or []),
              sets=sets if sets is not None else set_map(S), codes=[str(b.get("code")) for b in S.get("bom") or []])
    for r in S2.get("bom") or []:
        for k in ("name", "spec", "note"):
            if isinstance(r.get(k), str) and r[k]:
                r[k] = clean(r[k], ctx)
    rem = "\n".join(clean(x, ctx) if x.strip() else x for x in str(remarks or "").split("\n"))
    return S2, rem, ctx


__all__ = ["RULES_PATH", "CONTACT_FIELDS", "ASK_TITLE", "load_rules", "radius_c", "raw_radii", "round_radius", "spec_diam",
           "slide_move_tags", "gate_moves",
           "contact_of", "contact_line", "contact_manager", "name_map", "set_map", "Ctx", "clean", "scan_text",
           "pptx_items", "xlsx_items", "pptx", "xlsx_summary"]
